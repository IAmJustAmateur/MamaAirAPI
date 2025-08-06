from django.core.management.base import BaseCommand
from openpyxl import Workbook
from api.models import (
    RiskDefinitionMommySymptom,
    RiskDefinitionBabySymptom,
    UserRiskFactor,
)


class Command(BaseCommand):
    help = "Export risk-symptom links and user risk factors to an Excel file."

    def handle(self, *args, **options):
        wb = Workbook()

        # --- Sheet 1: RiskDefinitionMommySymptom ---
        ws_mommy = wb.active
        ws_mommy.title = "Mommy Symptoms"
        ws_mommy.append(["Risk", "Symptom"])

        mommy_links = RiskDefinitionMommySymptom.objects.select_related(
            "risk_definition", "symptom"
        ).order_by("risk_definition__name", "symptom__name")

        for link in mommy_links:
            ws_mommy.append([link.risk_definition.name, link.symptom.name])

        # --- Sheet 2: RiskDefinitionBabySymptom ---
        ws_baby = wb.create_sheet("Baby Symptoms")
        ws_baby.append(["Risk", "Symptom"])

        baby_links = RiskDefinitionBabySymptom.objects.select_related(
            "risk_definition", "symptom"
        ).order_by("risk_definition__name", "symptom__name")

        for link in baby_links:
            ws_baby.append([link.risk_definition.name, link.symptom.name])

        # --- Sheet 3: UserRiskFactor ---
        ws_factors = wb.create_sheet("User Risk Factors")
        ws_factors.append(["Risk", "Condition", "Multiplier"])

        factors = UserRiskFactor.objects.select_related("risk").order_by("risk__name")
        for factor in factors:
            ws_factors.append([factor.risk.name, factor.condition, factor.multiplier])

        # Save Excel file
        file_name = "risk_export.xlsx"
        wb.save(file_name)

        self.stdout.write(self.style.SUCCESS(f"Export completed: {file_name}"))
